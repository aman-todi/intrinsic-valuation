"""Excel builder unit tests (no LibreOffice): defined names, formula hygiene, cached values.

The cached value written with every formula comes from the expression DSL twin, so these
tests also pin the twin math to the real engine (value per share, scenarios, full grid).
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.export.excel.builder import build_workbook
from app.export.excel.recalc_verify import read_defined_names
from tests.unit.excel_cases import CASES, ExcelCase, case

OUTPUT_NAMES = ("operating_value", "enterprise_value", "equity_value", "value_per_share", "upside_pct")
ALLOWED_LITERALS = {0.0, 1.0}
_STRING = re.compile(r'"(?:[^"]|"")*"')
_NUMBER = re.compile(r"(?<![A-Za-z_$\d.])(\d+(?:\.\d+)?(?:[eE][-+]?\d+)?)")


@pytest.fixture(scope="module")
def built(tmp_path_factory) -> dict[str, tuple[ExcelCase, object, Path]]:
    out: dict[str, tuple[ExcelCase, object, Path]] = {}
    root = tmp_path_factory.mktemp("xlsx")
    for cid in CASES:
        c = case(cid)
        result = c.compute()
        path = root / f"{cid}.xlsx"
        build_workbook(
            result, c.financials, c.market, c.assumptions, path, company_name="Test Co", model_reason="r"
        )
        out[cid] = (c, result, path)
    return out


def assumption_names(c: ExcelCase) -> list[str]:
    a = c.assumptions
    if c.model_type != "sotp":
        return [name for name, _ in a]
    names = ["corporate_overhead_capitalized"]
    for i, seg in enumerate(a.segments, start=1):
        names.append(f"seg{i}_ev_ebitda_multiple")
        if seg.segment_ebitda_margin is not None:
            names.append(f"seg{i}_segment_ebitda_margin")
        if seg.fcff_assumptions is not None:
            names += [f"seg{i}_{name}" for name, _ in seg.fcff_assumptions]
    if a.consolidated_fcff is not None:
        names += [f"cons_{name}" for name, _ in a.consolidated_fcff]
    return names


def assumption_values(c: ExcelCase) -> set[float]:
    a = c.assumptions
    if c.model_type != "sotp":
        return {f.value for _, f in a}
    vals = {a.corporate_overhead_capitalized.value}
    for seg in a.segments:
        vals.add(seg.ev_ebitda_multiple)
        if seg.segment_ebitda_margin is not None:
            vals.add(seg.segment_ebitda_margin)
        if seg.fcff_assumptions is not None:
            vals |= {f.value for _, f in seg.fcff_assumptions}
    if a.consolidated_fcff is not None:
        vals |= {f.value for _, f in a.consolidated_fcff}
    return vals


def formulas(path: Path) -> dict[tuple[str, str], str]:
    wb = load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    out[(ws.title, cell.coordinate)] = cell.value
    return out


def literals(formula: str) -> list[float]:
    body = _STRING.sub("", formula)
    body = re.sub(r"[A-Za-z_]+!", "", body)  # sheet qualifiers
    return [float(m) for m in _NUMBER.findall(body)]


def test_names_defined_and_formulas_have_no_raw_literals(built):
    """Every model type: each assumption/output has a defined name, and no formula hard-codes a
    literal (only 0/1) or a raw Assumptions!cell address, so edits to a named input flow through."""
    for cid, (c, _, path) in built.items():
        _check_workbook_hygiene(cid, c, path)


def _check_workbook_hygiene(cid, c, path):
    names = set(load_workbook(path).defined_names.keys())
    missing = [n for n in [*assumption_names(c), *OUTPUT_NAMES] if n not in names]
    assert not missing, (cid, missing)
    for label in ("base", "bull", "bear"):
        assert f"scenario_{label}_value_per_share" in names
    fs = formulas(path)
    assert fs
    bad = {k: f for k, f in fs.items() if not set(literals(f)) <= ALLOWED_LITERALS}
    assert not bad, list(bad.items())[:5]
    distinctive = {v for v in assumption_values(c) if v not in ALLOWED_LITERALS}
    for f in fs.values():
        assert not distinctive & set(literals(f))
    # the projection and valuation sheets are driven by defined names, not raw cell addresses
    for sheet in ("Projections", "Valuation"):
        sheet_formulas = [f for (s, _), f in fs.items() if s == sheet]
        used = {
            n for n in names for f in sheet_formulas if re.search(rf"(?<![\w.]){re.escape(n)}(?![\w(])", f)
        }
        assert len(used) >= 3, (sheet, used)
    assert not any("Assumptions!" in f for f in fs.values())  # never a raw Assumptions cell address


def test_cached_values_match_engine(built):
    """The cached value written with every formula (from the expression-DSL twin) equals the real
    engine for every model type: value per share, bridge, scenarios and the full grid."""
    for cid, (_, result, path) in built.items():
        _check_cached_values(cid, result, path)


def _check_cached_values(cid, result, path):
    names = ["value_per_share", "equity_value", "enterprise_value", "operating_value", "upside_pct"]
    names += [
        "scenario_base_value_per_share",
        "scenario_bull_value_per_share",
        "scenario_bear_value_per_share",
    ]
    grid = result.sensitivity_grid
    if len(grid) == 25:
        names += [f"sens_row{k // 5}_col{k % 5}" for k in range(25)]
    got = read_defined_names(path, names)
    expected = [
        result.value_per_share,
        result.equity_value,
        result.enterprise_value,
        result.operating_value,
        result.upside_pct,
        *(s.value_per_share for s in result.scenarios),
    ]
    if len(grid) == 25:
        expected += [cell.value_per_share for cell in grid]
    for name, exp in zip(names, expected, strict=True):
        assert got[name] == pytest.approx(exp, rel=1e-9, abs=1e-9), (cid, name)
