"""Builds real workbooks from REAL engine results and recalculates them in headless
LibreOffice (spec §13.1 test_excel_recalc_verification).

Skips with a clear message when ``soffice`` is missing or cannot convert a spreadsheet
(e.g. the ``libreoffice-calc`` component is not installed: "source file could not be loaded").
"""

from __future__ import annotations

import asyncio
import tempfile
from pathlib import Path

import pytest
import xlsxwriter
from openpyxl import load_workbook

from app.export.excel.builder import build_workbook
from app.export.excel.recalc_verify import (
    ExcelRecalcError,
    ExcelVerificationError,
    find_soffice,
    recalc_and_read,
    relative_error,
    verify_workbook,
)
from app.valuation import get_valuator
from tests.unit.excel_cases import CASES, ExcelCase, case

pytestmark = pytest.mark.soffice

TOL = 0.005
SENS_CELL = "sens_row1_col3"

_probe_result: list[str | None] = []


def _probe() -> str | None:
    """None when soffice can recalculate a workbook, else the skip reason (cached)."""
    if _probe_result:
        return _probe_result[0]
    reason: str | None = None
    if find_soffice() is None:
        reason = "LibreOffice (soffice) not on PATH; install libreoffice-calc to run Excel recalc tests"
    else:
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "probe.xlsx"
            wb = xlsxwriter.Workbook(str(path))
            ws = wb.add_worksheet("S")
            ws.write("A1", 2)
            wb.define_name("probe_in", "=S!$A$1")
            ws.write_formula("B1", "=probe_in*3", None, 0)  # stale cached value: must be recalculated
            wb.define_name("probe_out", "=S!$B$1")
            wb.close()
            try:
                got = asyncio.run(recalc_and_read(path, ["probe_out"], timeout=120))["probe_out"]
                if got != 6:
                    reason = f"LibreOffice did not recalculate on load (got {got}, expected 6)"
            except ExcelRecalcError as exc:
                reason = f"LibreOffice cannot convert spreadsheets (is libreoffice-calc installed?): {exc}"
    _probe_result.append(reason)
    return reason


@pytest.fixture(scope="module", autouse=True)
def _needs_soffice():
    reason = _probe()
    if reason:
        pytest.skip(reason)


def _build(c: ExcelCase, tmp_path: Path):
    result = c.compute()
    path = tmp_path / f"{c.id}.xlsx"
    build_workbook(result, c.financials, c.market, c.assumptions, path, company_name="Test Co")
    return result, path


def _grid_cell(result, name: str) -> float:
    i, j = (int(x) for x in name.removeprefix("sens_row").split("_col"))
    return result.sensitivity_grid[i * 5 + j].value_per_share


@pytest.mark.parametrize("cid", list(CASES))
async def test_recalculated_workbook_matches_engine(cid, tmp_path):
    c = case(cid)
    result, path = _build(c, tmp_path)
    names = [
        "value_per_share",
        "equity_value",
        "scenario_bull_value_per_share",
        "scenario_bear_value_per_share",
    ]
    if len(result.sensitivity_grid) == 25:
        names.append(SENS_CELL)
    got = await recalc_and_read(path, names, timeout=120)

    assert relative_error(got["value_per_share"], result.value_per_share) <= TOL
    assert got["equity_value"] == pytest.approx(result.equity_value, rel=TOL)
    bull = next(s for s in result.scenarios if s.label == "bull")
    bear = next(s for s in result.scenarios if s.label == "bear")
    assert relative_error(got["scenario_bull_value_per_share"], bull.value_per_share) <= TOL
    assert relative_error(got["scenario_bear_value_per_share"], bear.value_per_share) <= TOL
    if SENS_CELL in names:
        assert relative_error(got[SENS_CELL], _grid_cell(result, SENS_CELL)) <= TOL


async def test_verify_workbook_passes_and_detects_mismatch(tmp_path):
    c = case("fcff_mature")
    result, path = _build(c, tmp_path)
    values = await verify_workbook(path, result, timeout=120)
    assert values["value_per_share"] == pytest.approx(result.value_per_share, rel=TOL)

    wrong = result.model_copy(update={"value_per_share": result.value_per_share * 1.02})
    with pytest.raises(ExcelVerificationError) as exc:
        await verify_workbook(path, wrong, timeout=120)
    assert exc.value.actual == pytest.approx(result.value_per_share, rel=TOL)


def _set_input(path: Path, out: Path, name: str, value: float) -> None:
    """Edit a named input cell with openpyxl (which also drops every cached formula value)."""
    wb = load_workbook(path)
    sheet, coord = next(wb.defined_names[name].destinations)
    wb[sheet][coord.replace("$", "")].value = value
    wb.save(out)


async def test_terminal_growth_edit_flows_through(tmp_path):
    """Editing a named input in the workbook changes value per share (and the grid centre) to what
    the engine computes for the edited assumptions."""
    cid = "fcff_mature"
    c = case(cid)
    result, path = _build(c, tmp_path)
    new_g = c.assumptions.terminal_growth_rate.value + 0.01
    edited = tmp_path / f"{cid}_edited.xlsx"
    _set_input(path, edited, "terminal_growth_rate", new_g)

    got = await recalc_and_read(edited, ["value_per_share", "sens_row2_col2"], timeout=120)

    edited_a = c.assumptions.model_copy(
        update={
            "terminal_growth_rate": c.assumptions.terminal_growth_rate.model_copy(update={"value": new_g})
        }
    )
    expected = get_valuator(c.model_type).compute(c.financials, c.market, edited_a, c.window).value_per_share
    assert expected > result.value_per_share  # higher terminal growth (ROIC > WACC here) raises value
    assert got["value_per_share"] > result.value_per_share
    assert relative_error(got["value_per_share"], expected) <= TOL
    assert relative_error(got["sens_row2_col2"], expected) <= TOL  # grid centre follows the edit too
